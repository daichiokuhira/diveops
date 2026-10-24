"""VBAの ImportAppCsv（ledger/make_bas.py）と同じ手順をPythonで再現する（テスト用）。
関数名・判定はマクロと1対1にしてある。片方だけ直さないこと。
使い方：python3 ledger/import_port.py 台帳.xlsx 取り込むCSV [CSV ...] 保存先.xlsx
"""
import sys, json, pathlib, datetime as dt
from openpyxl import load_workbook
ROOT = pathlib.Path(__file__).resolve().parents[1]
G = json.load(open(ROOT / 'ledger' / 'build' / 'grids.json', encoding='utf-8'))
LIM = G['limits']
D_FIRST, D_LAST, K_FIRST, K_LAST = LIM['D_FIRST'], LIM['D_LAST'], LIM['K_FIRST'], LIM['K_LAST']
MAP = [tuple(x) for x in G['import']]          # (CSVの見出し, 列, 種類)
INC = [(n, c, k) for c, n, k in G['incident']]  # (CSVの見出し, 列, 種類)
C_RECID = next(c for n, c, k in MAP if n == '記録ID')
C_TS = next(c for n, c, k in MAP if n == '記録更新日時')
C_NOTE = next(c for n, c, k in MAP if n == '修正記録')
NEXT_COLS = G['next_cols']
EPOCH = dt.datetime(1899, 12, 30)


class LimitError(Exception):
    pass


# ---- VBA の組み込み関数のまね（必要な範囲だけ） ----
def cstr(v):
    """CStr(cell.Value)。空は ""、整数の Double は小数点なし"""
    if v is None: return ''
    if isinstance(v, bool): return 'True' if v else 'False'
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    if isinstance(v, dt.datetime):
        return v.strftime('%Y/%m/%d') if v.time() == dt.time(0) else v.strftime('%Y/%m/%d %H:%M:%S')
    if isinstance(v, dt.time): return v.strftime('%H:%M:%S')
    return str(v)

def is_date(s):
    return to_date(s) is not None or to_time(s) is not None

def to_date(s):
    for f in ('%Y/%m/%d', '%Y-%m-%d', '%Y/%m/%d %H:%M', '%Y-%m-%d %H:%M'):
        try: return dt.datetime.strptime(s.strip(), f)
        except ValueError: pass
    return None

def to_time(s):
    try:
        p = s.strip().split(':')
        if len(p) in (2, 3) and all(x.isdigit() for x in p):
            h, m = int(p[0]), int(p[1])
            if h < 24 and m < 60: return dt.time(h, m)
    except Exception: pass
    return None

def is_numeric(s):
    try: float(s.strip()); return s.strip() != ''
    except ValueError: return False

def serial(v):
    """セルの日付・時刻を Excel の数値（日）に"""
    if isinstance(v, dt.datetime): return (v - EPOCH).total_seconds() / 86400
    if isinstance(v, dt.time): return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if isinstance(v, dt.date): return (dt.datetime(v.year, v.month, v.day) - EPOCH).days
    return float(v)

def vba_round(x): return round(x)  # VBA の Round も偶数丸め


# ---- CSV ----
def parse_csv(s):
    s = s.replace('﻿', ''); res = []; fields = []; cur = ''; inq = False; i = 0
    while i < len(s):
        ch = s[i]
        if inq:
            if ch == '"':
                if i + 1 < len(s) and s[i + 1] == '"': cur += '"'; i += 1
                else: inq = False
            else: cur += ch
        else:
            if ch == '"': inq = True
            elif ch == ',': fields.append(cur); cur = ''
            elif ch == '\r': pass
            elif ch == '\n': fields.append(cur); cur = ''; res.append(fields); fields = []
            else: cur += ch
        i += 1
    if cur or fields: fields.append(cur); res.append(fields)
    return res

def read_text(path, charset):
    # ADODB.Stream と同じく、読めないバイトは置き換えて読み進める
    return open(path, 'rb').read().decode('utf-8' if charset == 'UTF-8' else 'cp932', errors='replace')

def header_map(rows):
    hdr = {}
    if rows:
        for c, h in enumerate(rows[0]):
            nm = h.strip()
            if nm not in hdr: hdr[nm] = c
        if '減圧停止' in hdr and '減圧停止(表の判定)' not in hdr: hdr['減圧停止(表の判定)'] = hdr['減圧停止']
    return hdr

def is_app_header(hdr): return '事案ID' in hdr and '職員番号' in hdr and '本数' in hdr

def fld(v, hdr, name, kind):
    s = ''
    if name in hdr and hdr[name] < len(v): s = v[hdr[name]]
    if len(s) >= 2 and s[0] == "'" and s[1] in '=+-@\t': s = s[1:]   # 列の種類は問わない
    return s


# ---- セル ----
def cell_str(c): return cstr(c.value)

def put_val(c, s, kind):
    if s == '': c.value = None; return
    if kind == 'd':
        d = to_date(s)
        if d is not None: c.number_format = 'yyyy/m/d'; c.value = d
        else: c.number_format = '@'; c.value = s
    elif kind == 'h':
        t = to_time(s)
        if t is not None: c.number_format = 'h:mm'; c.value = t
        else: c.number_format = '@'; c.value = s
    elif kind == 'n':
        if is_numeric(s):
            c.number_format = 'General'; x = float(s); c.value = int(x) if x.is_integer() and '.' not in s else x
        else: c.number_format = '@'; c.value = s
    else:
        c.number_format = '@'; c.value = s
    if isinstance(c.value, str): c.data_type = 's'   # 「=」で始まっても数式にしない（Excel の文字の書式「@」と同じ）

def same_val(c, s, kind):
    v = c.value
    if isinstance(v, str) and v.startswith('#'): return False  # エラー値
    if s == '': return cstr(v) == ''
    if v is None: return False
    if kind == 'd':
        d = to_date(s)
        if d is not None and isinstance(v, (dt.datetime, dt.date)): return serial(d) == serial(v)
        return cstr(v) == s
    if kind == 'h':
        t = to_time(s)
        if t is not None and isinstance(v, (dt.datetime, dt.time, int, float)) and not isinstance(v, bool):
            x = serial(v); return vba_round((x - int(x)) * 1440) == vba_round(serial(t) * 1440)
        return cstr(v) == s
    if kind == 'n':
        if is_numeric(s) and isinstance(v, (int, float)) and not isinstance(v, bool): return float(s) == float(v)
        return cstr(v) == s
    return cstr(v) == s

def last_data_row(ws, first, last, cols):
    top = first - 1
    for c in cols:
        for r in range(last, first - 1, -1):
            if ws.cell(r, c).value not in (None, ''):
                top = max(top, r); break
    return top

def next_empty_row(ws, first, last, cols):
    r = last_data_row(ws, first, last, cols)
    return 0 if r >= last else r + 1


# ---- 照合の索引 ----
def row_key(wd, r): return f"{cell_str(wd.cell(r, 19))}|{cell_str(wd.cell(r, 2))}|{cell_str(wd.cell(r, 20))}"

def reindex(wd, r, ids, keys_all, keys_noid):
    rid = cell_str(wd.cell(r, C_RECID))
    if rid != '': ids.setdefault(rid, r)
    if cell_str(wd.cell(r, 19)) != '':
        key = row_key(wd, r); keys_all.setdefault(key, r)
        if rid == '': keys_noid.setdefault(key, r)

def unindex(wd, r, ids, keys_all, keys_noid):
    rid = cell_str(wd.cell(r, C_RECID)); key = row_key(wd, r)
    if rid != '' and ids.get(rid) == r: del ids[rid]
    if keys_all.get(key) == r: del keys_all[key]
    if keys_noid.get(key) == r: del keys_noid[key]

def is_newer(new_ts, old_ts):
    """IsNewer：CSVの記録更新日時が台帳より新しいか（文字で比べる）"""
    if old_ts == '': return True
    if new_ts == '': return False
    return new_ts > old_ts

def find_row(rid, key, ids, keys_all, keys_noid):
    if rid != '':
        if rid in ids: return ids[rid]
        if key in keys_noid: return keys_noid[key]
        return 0
    return keys_all.get(key, 0)

def write_row(wd, r, v, hdr, is_new):
    for name, col, kind in MAP:
        if is_new or name in hdr: put_val(wd.cell(r, col), fld(v, hdr, name, kind), kind)

def row_same(wd, r, v, hdr):
    for name, col, kind in MAP:
        if name in hdr and not same_val(wd.cell(r, col), fld(v, hdr, name, kind), kind): return False
    return True

def clear_row(wd, r, rid, ts):
    for name, col, kind in MAP: wd.cell(r, col).value = None
    if rid != '':
        put_val(wd.cell(r, C_RECID), rid, 't'); put_val(wd.cell(r, C_TS), ts, 't')
        put_val(wd.cell(r, C_NOTE), '削除済み（アプリで削除）', 't')


# ---- 本体 ----
def new_counts(): return dict(added=0, updated=0, deleted=0, skipped=0, inc_added=0, bad=0, notes=[])

def import_one(wb, path, n, now='取り込み試験'):
    """ImportOne。n は件数の dict（足し込む）。取り込まなかったときは理由の文字列を返す"""
    rows = parse_csv(read_text(path, 'UTF-8'))
    hdr = header_map(rows)
    name = pathlib.Path(path).name
    if not is_app_header(hdr):
        if is_app_header(header_map(parse_csv(read_text(path, 'Shift_JIS')))):
            return f'{name}：Excelで保存し直したCSVです。メールで届いた元のファイルを使ってください。'
        return f'{name}：アプリのCSVではありません（見出しに 事案ID・職員番号・本数 がありません）。'
    ncol = len(rows[0])
    wd, wk = wb['潜水記録'], wb['事案']
    ids, keys_all, keys_noid = {}, {}, {}
    for r in range(D_FIRST, D_LAST + 1): reindex(wd, r, ids, keys_all, keys_noid)
    incs = {cell_str(wk.cell(r, 1)) for r in range(K_FIRST, K_LAST + 1) if cell_str(wk.cell(r, 1)) != ''}
    next_d = next_empty_row(wd, D_FIRST, D_LAST, NEXT_COLS)
    next_k = next_empty_row(wk, K_FIRST, K_LAST, [1, 2])
    for v in rows[1:]:
        if len(v) < ncol:
            if ''.join(v).strip() != '': n['bad'] += 1
            continue
        iid = fld(v, hdr, '事案ID', 't'); rid = fld(v, hdr, '記録ID', 't'); st = fld(v, hdr, '状態', 't')
        ts = fld(v, hdr, '記録更新日時', 't')
        key = f"{iid}|{fld(v, hdr, '職員番号', 't')}|{fld(v, hdr, '本数', 't')}"
        found = find_row(rid, key, ids, keys_all, keys_noid)
        live = False
        if found > 0 and not is_newer(ts, cell_str(wd.cell(found, C_TS))):
            n['skipped'] += 1
            live = cell_str(wd.cell(found, 19)) != ''   # 削除の印の行なら、事案も足さない
        elif st == '削除':
            if found > 0:
                prior = cell_str(wd.cell(found, C_RECID))
                unindex(wd, found, ids, keys_all, keys_noid)
                if rid == '': rid = prior
                clear_row(wd, found, rid, ts); reindex(wd, found, ids, keys_all, keys_noid); n['deleted'] += 1
            elif rid != '':
                # まだ台帳にない記録の削除：削除の印の行を作る
                if next_d == 0: raise LimitError(f'潜水記録の上限（{D_LAST - D_FIRST + 1}件）に達しました。')
                clear_row(wd, next_d, rid, ts); reindex(wd, next_d, ids, keys_all, keys_noid); n['deleted'] += 1
                next_d = 0 if next_d >= D_LAST else next_d + 1
        elif found > 0:
            live = True
            if row_same(wd, found, v, hdr): n['skipped'] += 1
            else:
                unindex(wd, found, ids, keys_all, keys_noid); write_row(wd, found, v, hdr, False)
                reindex(wd, found, ids, keys_all, keys_noid); n['updated'] += 1
        else:
            if next_d == 0: raise LimitError(f'潜水記録の上限（{D_LAST - D_FIRST + 1}件）に達しました。')
            write_row(wd, next_d, v, hdr, True); reindex(wd, next_d, ids, keys_all, keys_noid); n['added'] += 1; live = True
            next_d = 0 if next_d >= D_LAST else next_d + 1
        if live and iid != '' and iid not in incs:
            if next_k == 0: raise LimitError(f'事案の上限（{K_LAST - K_FIRST + 1}件）に達しました。')
            for nm, col, kind in INC: put_val(wk.cell(next_k, col), fld(v, hdr, nm, kind), kind)
            put_val(wk.cell(next_k, 16), now, 't')
            incs.add(iid); n['inc_added'] += 1
            next_k = 0 if next_k >= K_LAST else next_k + 1
    return ''

def import_csv(wb, paths):
    """ImportAppCsv（ファイルを選んだ後の部分）"""
    n = new_counts()
    for p in paths:
        note = import_one(wb, p, n)
        if note: n['notes'].append(note)
    return n


if __name__ == '__main__':
    src, *csvs, out = sys.argv[1:]
    wb = load_workbook(src)
    n = import_csv(wb, csvs)
    wb.save(out)
    print({k: v for k, v in n.items()})

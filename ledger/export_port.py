# VBAの「アプリ用に書き出す」（ExportTableJson）と同じ手順をPythonで再現し、結果をアプリの点検に通す
# 表の見出し行は grids.json（build_ledger.py が書く）から読む。値の読み方は VBA の JNum／HmToMin／JStr に合わせる
import json, sys, re, pathlib, datetime as dt
from openpyxl import load_workbook
ROOT = pathlib.Path(__file__).resolve().parents[1]
GR = json.load(open(ROOT / 'ledger' / 'build' / 'grids.json', encoding='utf-8'))['grids']
EPOCH = dt.datetime(1899, 12, 30)

def cstr(v):
    """CStr(v)"""
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    if isinstance(v, dt.datetime): return v.strftime('%Y/%m/%d') if v.time() == dt.time(0) else v.strftime('%Y/%m/%d %H:%M:%S')
    if isinstance(v, dt.time): return v.strftime('%H:%M:%S')
    return str(v)

def val(s):
    """VBA の Val：先頭から読める数字だけ（空白は無視）。読めなければ 0"""
    m = re.match(r'[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?', s.replace(' ', '').replace('\t', ''))
    if not m: return 0
    x = float(m.group(0)); return int(x) if x.is_integer() else x

def jnum(v): return val(cstr(v))          # JNum（空は 0）
def jstr(v): return cstr(v)               # JStr（空は ""）

def serial(v):
    if isinstance(v, dt.datetime): return (v - EPOCH).total_seconds() / 86400
    if isinstance(v, dt.time): return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if isinstance(v, dt.timedelta): return v.total_seconds() / 86400
    return float(v)

def hm_to_min(v):
    """HmToMin：時刻型・数値・「0:10」・コロンのない「10」（分）"""
    if isinstance(v, (dt.datetime, dt.time, dt.timedelta)): return round(serial(v) * 1440)
    if isinstance(v, (int, float)) and not isinstance(v, bool): return round(v * 1440) if v < 1 else round(v)
    s = str(v).strip().replace('：', ':'); p = s.split(':')
    return int(val(p[0])) * 60 + int(val(p[1])) if len(p) >= 2 else int(val(s))

def main(src, out):
    ws = load_workbook(src)['潜水表']
    raw = lambda r, c: ws.cell(r, c).value
    def cell(r, c):
        v = raw(r, c)
        return None if cstr(v).strip() == '' else v
    def keys_row(r):
        out = []; c = 2
        while cell(r, c) is not None: out.append(cell(r, c)); c += 1
        return out
    def keys_col(r):
        out = []; rr = r + 1
        while cell(rr, 1) is not None: out.append(cell(rr, 1)); rr += 1
        return out
    b8 = raw(8, 2)
    t = {"format": "dive-table/1", "name": jstr(raw(3, 2)), "version": jstr(raw(4, 2)), "source": jstr(raw(5, 2)),
         "verified": cstr(raw(6, 2)) == "照合済み", "verifiedBy": jstr(raw(7, 2)),
         "updatedAt": b8.strftime('%Y-%m-%d') if isinstance(b8, dt.datetime) else cstr(b8),
         "notes": [x for x in cstr(raw(9, 2)).split("／") if x.strip() != '']}
    rk = ["minSurfaceInterval", "clearAfterMinutes", "avoidRepetitiveDeeperThan", "safetyStopMinutes", "safetyStopDepth", "maxAscentRate"]
    t["rules"] = {k: jnum(raw(3 + i, 13)) for i, k in enumerate(rk)}
    MDT, G1, DC, RN, AM, SL, SH = (GR[k][0] for k in ("MDT", "表1 グループの時間（分）", "表1 減圧停止（分）", "表3 残留窒素時間（RNT・分）",
                                                      "表3 修正最大潜水時間（AMDT・分）", "表2 水面休息 下限（時:分）", "表2 水面休息 上限（時:分）"))
    depths = keys_row(MDT); t["depths"] = [jnum(d) for d in depths]
    t["mdt"] = {str(jnum(d)): jnum(raw(MDT + 1, 2 + j)) for j, d in enumerate(depths)}
    letters = keys_row(G1)
    t["groups"] = {}; t["deco"] = {}
    for i, d in enumerate(keys_col(G1)):
        cuts = []
        for j, g in enumerate(letters):
            v = cell(G1 + 1 + i, 2 + j)
            if v is not None:
                cuts.append([jnum(v), jstr(g)]); s = cell(DC + 1 + i, 2 + j)
                if s is not None: t["deco"][f"{jnum(d)}-{cstr(g)}"] = {"limit": jnum(v), "stop5": jnum(s)}
        t["groups"][str(jnum(d))] = cuts
    for key, top in (("rnt", RN), ("amdt", AM)):
        cols = [str(val(cstr(x))) for x in keys_row(top)]; t[key] = {}
        for i, g in enumerate(keys_col(top)):
            row = {cols[j]: jnum(cell(top + 1 + i, 2 + j)) for j in range(len(cols)) if cell(top + 1 + i, 2 + j) is not None}
            if row: t[key][cstr(g)] = row
    ends = [cstr(x).replace('終了', '') for x in keys_row(SL)]; news = keys_col(SL)
    t["surfaceInterval"] = {}
    for j, eg in enumerate(ends):
        rows = []
        for i in range(len(news) - 1, -1, -1):
            lo, hi = cell(SL + 1 + i, 2 + j), cell(SH + 1 + i, 2 + j)
            if lo is not None and hi is not None: rows.append([hm_to_min(lo), hm_to_min(hi), cstr(news[i])])
        if rows: t["surfaceInterval"][eg] = rows
    json.dump(t, open(out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    return t

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])

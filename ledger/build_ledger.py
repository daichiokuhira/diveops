"""潜水管理台帳 第2版（Excel 2019・モノクロ・健診なし）"""
import json, datetime as dt, sys, pathlib
ROOT = pathlib.Path(__file__).resolve().parents[1]
BUILD = ROOT/"ledger"/"build"; BUILD.mkdir(exist_ok=True)
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as CL

OUT = sys.argv[1]
TABLE = json.load(open(ROOT/'app'/'tables'/'naui-2022.json'))
SAMPLE = json.load(open(BUILD/'sample_rows.json'))
FONT = "游ゴシック"
BK = "000000"
def f(size=10.5, bold=False): return Font(name=FONT, size=size, bold=bold, color=BK)
hair = Side(style="hair", color=BK); thin = Side(style="thin", color=BK); med = Side(style="medium", color=BK); thick = Side(style="thick", color=BK)
DATE = "yyyy/m/d"; TIME = "h:mm"
BOLD = Font(name=FONT, bold=True, color=BK)

M_FIRST, M_LAST = 5, 204
Q_FIRST, Q_LAST = 5, 1004
D_FIRST, D_LAST = 5, 3004
K_FIRST, K_LAST = 5, 1004
rng = lambda sh, c, a, b: f"'{sh}'!${c}${a}:${c}${b}"
M = lambda c: rng("隊員マスタ", c, M_FIRST, M_LAST)
Q = lambda c: rng("資格研修記録", c, Q_FIRST, Q_LAST)
D = lambda c: rng("潜水記録", c, D_FIRST, D_LAST)
K = lambda c: rng("事案", c, K_FIRST, K_LAST)

wb = Workbook()
wb.calculation.fullCalcOnLoad = True

def title(ws, text, sub=None, span=4):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(1, 1, text); c.font = f(18, True); c.alignment = Alignment(vertical="bottom")
    for col in range(1, span + 1): ws.cell(1, col).border = Border(bottom=thick)
    ws.row_dimensions[1].height = 34
    if sub: ws.cell(2, 1, sub).font = f(10)

def header(ws, row, names, widths=None, first=1):
    for i, n in enumerate(names):
        c = ws.cell(row, first + i, n); c.font = f(10, True)
        c.alignment = Alignment(wrap_text=True, vertical="bottom"); c.border = Border(bottom=med)
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths): ws.column_dimensions[CL(first + i)].width = w

def body(ws, r1, r2, c1, c2, center=()):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            x = ws.cell(r, c); x.font = f(); x.border = Border(bottom=hair)
            x.alignment = Alignment(vertical="center", horizontal="center" if c in center else None)

def input_box(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            x = ws.cell(r, c); x.font = f(); x.border = Border(left=hair, right=hair, top=hair, bottom=hair)

def dv_list(ws, name, cells):
    dv = DataValidation(type="list", formula1=f"={name}", allow_blank=True)
    dv.error = "プルダウンから選んでください（選択肢シートで追加できます）"; dv.errorTitle = "入力できない値です"
    ws.add_data_validation(dv); dv.add(cells)

def landscape(ws, portrait=False):
    ws.page_setup.orientation = "portrait" if portrait else "landscape"; ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True

def d(s): y, m, dd = map(int, s.split("/")); return dt.date(y, m, dd)

# ============ シートを用意（順番） ============
order = ["使い方", "隊員一覧", "資格一覧", "経歴簿", "報告書", "CSV抽出", "事案", "潜水記録", "隊員マスタ", "資格研修記録", "潜水表", "選択肢"]
ws0 = wb.active; ws0.title = order[0]
S = {n: (ws0 if n == order[0] else wb.create_sheet(n)) for n in order}

# ============ 選択肢 ============
wl = S["選択肢"]
title(wl, "選択肢", "ここを書き換えると、各シートのプルダウンに反映されます。", 4)
lists = {
    "A": ("所属", ["本署", "伊良部分署", "（分署名を入力）"]),
    "B": ("在籍状態", ["在籍", "異動", "退職"]),
    "C": ("潜水区分", ["災害", "訓練", "研修", "その他"]),
    "D": ("資格・研修名", ["潜水士免許", "水難救助研修", "小型船舶1級", "小型船舶2級", "小型船舶特殊", "救助技術研修"]),
    "E": ("資格種別", ["資格", "研修"]),
    "F": ("潜水地点", ["平良港", "佐良浜漁港", "渡口の浜", "訓練プール", "その他"]),
    "G": ("有無", ["有", "無"]),
}
header(wl, 4, [v[0] for v in lists.values()] + ["抽出用の区分"], [16, 12, 12, 18, 12, 16, 8, 14])
for col, (_, items) in lists.items():
    for i, it in enumerate(items): wl.cell(5 + i, ord(col) - 64, it).font = f()
wl.cell(5, 8, "すべて").font = f()
for i in range(10): wl.cell(6 + i, 8, f'=IF(C{5+i}="","",C{5+i})').font = f()
input_box(wl, 5, 20, 1, 7)
wl.cell(22, 1, "資格・研修名は先頭8件が「資格一覧」の列見出しになります。「潜水士免許」「水難救助研修」「小型船舶1級／2級／特殊」は判定に使うので名前を変えないでください。").font = f(9.5)
wl.freeze_panes = "A5"
for n, ref in {"L_所属": "$A$5:$A$20", "L_在籍": "$B$5:$B$20", "L_区分": "$C$5:$C$20", "L_資格名": "$D$5:$D$20",
               "L_種別": "$E$5:$E$20", "L_地点": "$F$5:$F$20", "L_有無": "$G$5:$G$20", "L_抽出区分": "$H$5:$H$15"}.items():
    wb.defined_names[n] = DefinedName(n, attr_text=f"'選択肢'!{ref}")
wb.defined_names["L_職員番号"] = DefinedName("L_職員番号", attr_text=M("A"))
wb.defined_names["L_事案ID"] = DefinedName("L_事案ID", attr_text=K("A"))

# ============ 隊員マスタ ============
wm = S["隊員マスタ"]
title(wm, "隊員マスタ", "1行に1人。職員番号は重複させないでください。異動・退職しても行は消さず、在籍状態を変えます。", 4)
header(wm, 4, ["職員番号", "氏名", "所属", "在籍状態", "備考"], [11, 16, 14, 11, 30])
members = [(1001, "見本 一郎", "本署", "在籍"), (1002, "見本 二郎", "本署", "在籍"), (1003, "見本 三郎", "伊良部分署", "在籍"),
           (1004, "見本 四郎", "伊良部分署", "在籍"), (1005, "見本 五郎", "本署", "在籍"), (1006, "見本 六郎", "伊良部分署", "在籍"),
           (1007, "見本 七郎", "本署", "異動"), (1008, "見本 八郎", "本署", "在籍")]
input_box(wm, M_FIRST, M_LAST, 1, 5)
for i, m in enumerate(members):
    for j, v in enumerate(m): wm.cell(M_FIRST + i, 1 + j, v)
wm.cell(M_FIRST, 5, "見本データ。使い始める前に削除")
wm.cell(M_FIRST + 7, 5, "潜水士免許なし（候補者）の例")
dv_list(wm, "L_所属", f"C{M_FIRST}:C{M_LAST}"); dv_list(wm, "L_在籍", f"D{M_FIRST}:D{M_LAST}")
dup = DataValidation(type="custom", formula1=f"=COUNTIF($A${M_FIRST}:$A${M_LAST},A{M_FIRST})=1"); dup.error = "同じ職員番号がすでにあります"; dup.errorTitle = "重複"
wm.add_data_validation(dup); dup.add(f"A{M_FIRST}:A{M_LAST}")
wm.freeze_panes = "A5"; wm.auto_filter.ref = f"A4:E{M_LAST}"

# ============ 資格研修記録 ============
wq = S["資格研修記録"]
title(wq, "資格・研修の記録", "1行に1件。取得・修了のたびに行を追加します。小型船舶は級ごとに1行（1級と特殊なら2行）。更新・進級したら行を追加し、有効期限を入れます。", 4)
header(wq, 4, ["職員番号", "氏名（自動）", "種別", "名称", "取得・修了日", "有効期限", "期別・証番号", "備考"], [11, 14, 9, 18, 14, 12, 16, 26])
quals = [(1001, "資格", "潜水士免許", "2012/6/1", None, "見本"), (1001, "研修", "水難救助研修", "2016/10/1", None, "第15期"),
         (1001, "資格", "小型船舶1級", "2022/4/1", "2027/3/31", None), (1001, "資格", "小型船舶特殊", "2022/4/1", "2027/3/31", "同じ免許証"),
         (1002, "資格", "潜水士免許", "2015/7/1", None, None), (1002, "研修", "水難救助研修", "2023/10/1", None, "第22期"),
         (1002, "資格", "小型船舶2級", "2017/6/1", "2022/5/31", "1級へ進級前"), (1002, "資格", "小型船舶1級", "2022/6/1", "2027/5/31", "進級"),
         (1003, "資格", "潜水士免許", "2018/5/1", None, None), (1004, "資格", "潜水士免許", "2019/9/1", None, None),
         (1004, "研修", "水難救助研修", "2025/10/1", None, "第24期"), (1004, "研修", "救助技術研修", "2021/6/1", None, None),
         (1004, "資格", "小型船舶2級", "2023/7/1", "2028/6/30", None), (1004, "資格", "小型船舶特殊", "2023/7/1", "2028/6/30", None),
         (1005, "資格", "潜水士免許", "2020/6/1", None, None), (1006, "資格", "潜水士免許", "2021/8/1", None, None),
         (1006, "資格", "小型船舶2級", "2020/5/1", "2025/4/30", "更新忘れの例"), (1007, "資格", "潜水士免許", "2014/7/1", None, None),
         (1008, "研修", "救助技術研修", "2025/6/1", None, None)]
input_box(wq, Q_FIRST, Q_LAST, 1, 8)
for i, (eid, kind, name, day, exp, memo) in enumerate(quals):
    r = Q_FIRST + i; wq.cell(r, 1, eid); wq.cell(r, 3, kind); wq.cell(r, 4, name); wq.cell(r, 5, d(day))
    if exp: wq.cell(r, 6, d(exp))
    if memo: wq.cell(r, 7, memo)
for r in range(Q_FIRST, Q_LAST + 1):
    wq.cell(r, 2, f'=IF(A{r}="","",IFERROR(INDEX({M("B")},MATCH(A{r},{M("A")},0)),"番号なし"))')
    wq.cell(r, 5).number_format = DATE; wq.cell(r, 6).number_format = DATE
dv_list(wq, "L_職員番号", f"A{Q_FIRST}:A{Q_LAST}"); dv_list(wq, "L_種別", f"C{Q_FIRST}:C{Q_LAST}"); dv_list(wq, "L_資格名", f"D{Q_FIRST}:D{Q_LAST}")
wq.freeze_panes = "A5"; wq.auto_filter.ref = f"A4:H{Q_LAST}"

# ============ 潜水記録 ============
wd = S["潜水記録"]
title(wd, "潜水記録", "1人が1本潜るごとに1行。アプリのCSVは［アプリのCSVを取り込む］で入ります（同じ記録は二重に入りません）。手で入力してもかまいません。", 5)
DCOLS = ["日付", "職員番号", "氏名（自動）", "区分", "事案・訓練名", "地点", "最大水深(m)", "潜水時間(分)", "前の潜水からの水面休息(分)", "タンク容量(L)",
         "開始圧(気圧)", "終了圧(気圧)", "水温(℃)", "透視度(m)", "減圧停止", "体調異常", "指揮者", "備考",
         "事案ID", "本数", "潜降時刻", "浮上時刻", "表の水深(m)", "グループ", "残留窒素時間(分)", "合計窒素時間(分)", "最大潜水時間(分)", "修正最大潜水時間(分)", "判定", "5m減圧停止(分)", "潜水表"]
DW = [11, 9, 12, 7, 20, 12, 8, 8, 10, 8, 8, 8, 7, 7, 7, 7, 11, 22, 13, 6, 7, 7, 8, 7, 9, 9, 9, 10, 10, 8, 22]
header(wd, 4, DCOLS, DW)
wd.cell(3, 19, "ここから右はアプリが計算した値").font = f(9, True)
HELP = ["年度（自動）", "経歴簿用（自動）", "抽出対象（自動）", "抽出用（自動）", "報告書用（自動）", "取込時の氏名"]
header(wd, 4, HELP, [9] * 6, first=32)  # AF..AK
X = "'CSV抽出'"
KEYS = ["日付", "職員番号", None, "区分", "事案名", "地点", "最大水深", "潜水時間", "水面休息", "タンク", "開始圧", "終了圧", "水温", "透視度", "減圧停止", "体調異常", "指揮者", "備考",
        "事案ID", "本数", "潜降", "浮上", "表の水深", "グループ", "RNT", "TNT", "MDT", "AMDT", "判定", "停止", "潜水表"]
def hmval(s): h, m = map(int, s.split(":")); return dt.time(h, m)
for i, row in enumerate(SAMPLE):
    r = D_FIRST + i
    for j, k in enumerate(KEYS):
        if not k: continue
        v = row[k]
        if v == "" or v is None: continue
        if k == "日付": v = d(v)
        if k in ("潜降", "浮上"): v = hmval(v)
        wd.cell(r, 1 + j, v)
    wd.cell(r, 37, row["氏名"])
input_box(wd, D_FIRST, D_LAST, 1, 31)
for r in range(D_FIRST, D_LAST + 1):
    wd.cell(r, 1).number_format = DATE; wd.cell(r, 21).number_format = TIME; wd.cell(r, 22).number_format = TIME
    wd.cell(r, 3, f'=IF(B{r}="","",IFERROR(INDEX({M("B")},MATCH(B{r},{M("A")},0)),IF(AK{r}="","番号なし",AK{r})))')
    wd.cell(r, 32, f'=IF(A{r}="","",YEAR(EDATE(A{r},-3)))')
    wd.cell(r, 33, f"=IF(OR(B{r}=\"\",'経歴簿'!$C$4=\"\"),\"\",IF(B{r}='経歴簿'!$C$4,COUNTIF($B${D_FIRST}:B{r},'経歴簿'!$C$4),\"\"))")
    wd.cell(r, 34, (f'=IF(A{r}="",0,IF(AND(OR({X}!$C$4="",A{r}>={X}!$C$4),OR({X}!$C$5="",A{r}<={X}!$C$5),'
                    f'OR({X}!$C$6="すべて",{X}!$C$6="",D{r}={X}!$C$6),OR({X}!$C$7="",B{r}={X}!$C$7)),1,0))'))
    wd.cell(r, 35, f'=IF(AH{r}=1,SUM($AH${D_FIRST}:AH{r}),"")')
    wd.cell(r, 36, f"=IF(OR(S{r}=\"\",'報告書'!$C$4=\"\"),\"\",IF(S{r}='報告書'!$C$4,COUNTIF($S${D_FIRST}:S{r},'報告書'!$C$4),\"\"))")
    for c in range(32, 38): wd.cell(r, c).font = f(9)
dv_list(wd, "L_職員番号", f"B{D_FIRST}:B{D_LAST}"); dv_list(wd, "L_区分", f"D{D_FIRST}:D{D_LAST}")
dv_list(wd, "L_地点", f"F{D_FIRST}:F{D_LAST}"); dv_list(wd, "L_有無", f"O{D_FIRST}:P{D_LAST}")
for c in range(32, 38):
    wd.column_dimensions[CL(c)].outlineLevel = 1; wd.column_dimensions[CL(c)].hidden = True
wd.freeze_panes = "C5"; wd.auto_filter.ref = f"A4:AE{D_LAST}"

# ============ 事案 ============
wk = S["事案"]
title(wk, "事案", "1事案1行。アプリのCSVを取り込むと自動で追加されます。手で追加するときは、事案IDを重ならないように付けてください（例：20261024-01）。", 5)
KC = ["事案ID", "日付", "区分", "事案名", "地点", "指揮者", "開始", "終了", "水温(℃)", "透視度(m)", "海況", "備考", "潜水表", "潜水員（自動）", "本数（自動）", "取り込み日時"]
header(wk, 4, KC, [13, 11, 7, 22, 14, 11, 15, 15, 7, 8, 10, 26, 22, 8, 8, 15])
incs = {}
for row in SAMPLE: incs.setdefault(row["事案ID"], row)
input_box(wk, K_FIRST, K_LAST, 1, 13)
for i, (iid, row) in enumerate(incs.items()):
    r = K_FIRST + i
    vals = [iid, d(row["日付"]), row["区分"], row["事案名"], row["地点"], row["指揮者"], row["事案開始"], row["事案終了"], row["水温"], row["透視度"], row["海況"], row["事案備考"], row["潜水表"]]
    for j, v in enumerate(vals):
        if v != "": wk.cell(r, 1 + j, v)
    wk.cell(r, 16, "見本")
for r in range(K_FIRST, K_LAST + 1):
    wk.cell(r, 2).number_format = DATE
    wk.cell(r, 14, f'=IF(A{r}="","",COUNTIFS({D("S")},A{r},{D("T")},1))')
    wk.cell(r, 15, f'=IF(A{r}="","",COUNTIF({D("S")},A{r}))')
    for c in (14, 15): wk.cell(r, c).font = f(); wk.cell(r, c).alignment = Alignment(horizontal="center")
dv_list(wk, "L_区分", f"C{K_FIRST}:C{K_LAST}")
kdup = DataValidation(type="custom", formula1=f"=COUNTIF($A${K_FIRST}:$A${K_LAST},A{K_FIRST})=1"); kdup.error = "同じ事案IDがすでにあります"; kdup.errorTitle = "重複"
wk.add_data_validation(kdup); kdup.add(f"A{K_FIRST}:A{K_LAST}")
wk.freeze_panes = "B5"; wk.auto_filter.ref = f"A4:P{K_LAST}"

# ============ 隊員一覧 ============
wv = S["隊員一覧"]
title(wv, "隊員一覧", None, 4)
wv.cell(2, 1, "基準日").font = f(10)
c = wv.cell(2, 2, "=TODAY()"); c.number_format = DATE; c.font = f(10, True)
wv.cell(2, 3, "年度").font = f(10)
wv.merge_cells("D2:E2")
c = wv.cell(2, 4, "=YEAR(EDATE(B2,-3))"); c.number_format = '0"年度"'; c.font = f(10, True); c.alignment = Alignment(horizontal="left")
wv.cell(2, 6, "このシートは自動表示です。入力は「隊員マスタ」「資格研修記録」「潜水記録」へ。").font = f(9.5)
H_ROW, V_FIRST = 7, 8; V_LAST = V_FIRST + (M_LAST - M_FIRST)
cols = ["職員番号", "氏名", "所属", "在籍", "潜水士免許", "水難救助研修", "小型船舶", "今年度 本数", "累計 本数", "累計 潜水時間(分)", "最大水深(m)", "最終潜水日", "最終潜水から(日)"]
header(wv, H_ROW, cols, [9, 12, 11, 6, 8, 9, 18, 8, 8, 10, 9, 11, 10])
vr = lambda col: f"${col}${V_FIRST}:${col}${V_LAST}"
G = vr("G")
kpis = [("A:B", "在籍隊員", f'=COUNTIF({vr("D")},"在籍")', "人"),
        ("C:E", "潜水士免許 保有", f'=COUNTIFS({vr("D")},"在籍",{vr("E")},"有")', "人"),
        ("F:G", "小型船舶 保有（有効）", f'=SUMPRODUCT(({vr("D")}="在籍")*((((LEFT({G},2)="1級")+(LEFT({G},2)="2級")+ISNUMBER(SEARCH("・特殊",{G}))+({G}="特殊"))>0)))', "人"),
        ("H:J", "今年度 潜水本数", f"=COUNTIF({D('AF')},D2)", "本")]
for span, lab, fm, unit in kpis:
    a, b = span.split(":"); wv.merge_cells(f"{a}4:{b}4"); wv.merge_cells(f"{a}5:{b}5")
    wv[f"{a}4"] = lab; wv[f"{a}4"].font = f(9.5)
    wv[f"{a}5"] = fm; wv[f"{a}5"].font = f(20, True); wv[f"{a}5"].number_format = f'0"{unit}"'; wv[f"{a}5"].alignment = Alignment(horizontal="left")
wv.row_dimensions[5].height = 30
def q_cnt(idref, name): return f'COUNTIFS({Q("A")},{idref},{Q("D")},"{name}")'
def q_valid(idref, name):
    b = f'{Q("A")},{idref},{Q("D")},"{name}"'
    return f'(COUNTIFS({b},{Q("F")},"")+COUNTIFS({b},{Q("F")},">="&TODAY()))'
def boat(idref):
    hull = (f'IF({q_valid(idref,"小型船舶1級")}>0,"1級",IF({q_valid(idref,"小型船舶2級")}>0,"2級",'
            f'IF({q_cnt(idref,"小型船舶1級")}+{q_cnt(idref,"小型船舶2級")}>0,"× 1・2級 期限切れ","")))')
    sp = f'IF({q_valid(idref,"小型船舶特殊")}>0,"特殊",IF({q_cnt(idref,"小型船舶特殊")}>0,"× 特殊 期限切れ",""))'
    return f'IF(({hull})&({sp})="","—",_xlfn.TEXTJOIN("・",TRUE,{hull},{sp}))'
for i in range(M_LAST - M_FIRST + 1):
    r, mr = V_FIRST + i, M_FIRST + i; A = f"A{r}"
    F = {1: f"=IF('隊員マスタ'!A{mr}=\"\",\"\",'隊員マスタ'!A{mr})", 2: f"=IF({A}=\"\",\"\",'隊員マスタ'!B{mr})",
         3: f"=IF({A}=\"\",\"\",'隊員マスタ'!C{mr})", 4: f"=IF({A}=\"\",\"\",'隊員マスタ'!D{mr})",
         5: f'=IF({A}="","",IF({q_cnt(A,"潜水士免許")}>0,"有","無"))',
         6: f'=IF({A}="","",IF({q_cnt(A,"水難救助研修")}>0,"修了","—"))',
         7: f'=IF({A}="","",{boat(A)})',
         8: f'=IF({A}="","",COUNTIFS({D("B")},{A},{D("AF")},$D$2))',
         9: f'=IF({A}="","",COUNTIF({D("B")},{A}))',
         10: f'=IF({A}="","",SUMIF({D("B")},{A},{D("H")}))',
         11: f'=IF({A}="","",IF(I{r}=0,"",_xlfn.MAXIFS({D("G")},{D("B")},{A})))',
         12: f'=IF({A}="","",IF(I{r}=0,"",_xlfn.MAXIFS({D("A")},{D("B")},{A})))',
         13: f'=IF(L{r}="","",$B$2-L{r})'}
    for c, fm in F.items(): wv.cell(r, c, fm)
    wv.cell(r, 12).number_format = DATE
body(wv, V_FIRST, V_LAST, 1, 13, center=(4, 5, 6, 8, 9, 10, 11, 13))
wv.conditional_formatting.add(f"E{V_FIRST}:E{V_LAST}", FormulaRule(formula=[f'E{V_FIRST}="無"'], font=BOLD))
wv.conditional_formatting.add(f"G{V_FIRST}:G{V_LAST}", FormulaRule(formula=[f'ISNUMBER(SEARCH("×",G{V_FIRST}))'], font=BOLD))
wv.cell(H_ROW, 7).comment = Comment("有効な級を表示します（1級と2級の両方の記録があれば1級）。特殊も持っていれば「1級・特殊」。期限切れは「×」付きで太字。", "台帳")
wv.freeze_panes = f"C{V_FIRST}"; wv.auto_filter.ref = f"A{H_ROW}:M{V_LAST}"
landscape(wv); wv.print_title_rows = f"{H_ROW}:{H_ROW}"

# ============ 資格一覧 ============
wx = S["資格一覧"]
title(wx, "資格・研修一覧", "○＝有効、× 期限切れ＝有効期限を過ぎた記録のみ、—＝記録なし、1級に含む＝1級が有効なので2級は不要。列見出しは「選択肢」の資格・研修名（先頭8件）。", 4)
header(wx, 4, ["職員番号", "氏名", "所属", "在籍"], [9, 12, 11, 6])
for k in range(8):
    c = wx.cell(4, 5 + k, f"=IF('選択肢'!D{5+k}=\"\",\"\",'選択肢'!D{5+k})"); c.font = f(10, True); c.border = Border(bottom=med)
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="bottom"); wx.column_dimensions[CL(5 + k)].width = 11
XL = 5 + (M_LAST - M_FIRST)
for i in range(M_LAST - M_FIRST + 1):
    r, mr = 5 + i, M_FIRST + i
    wx.cell(r, 1, f"=IF('隊員マスタ'!A{mr}=\"\",\"\",'隊員マスタ'!A{mr})")
    for c, col in ((2, "B"), (3, "C"), (4, "D")): wx.cell(r, c, f"=IF(A{r}=\"\",\"\",'隊員マスタ'!{col}{mr})")
    for k in range(8):
        col = CL(5 + k); base = f'{Q("A")},$A{r},{Q("D")},{col}$4'; v1 = q_valid(f"$A{r}", "小型船舶1級")
        wx.cell(r, 5 + k, (f'=IF(OR($A{r}="",{col}$4=""),"",IF(AND({col}$4="小型船舶2級",{v1}>0),"1級に含む",IF(COUNTIFS({base})=0,"—",'
                           f'IF(COUNTIFS({base},{Q("F")},"")+COUNTIFS({base},{Q("F")},">="&TODAY())>0,"○","× 期限切れ"))))'))
body(wx, 5, XL, 1, 12, center=tuple(range(4, 13)))
wx.conditional_formatting.add(f"E5:L{XL}", FormulaRule(formula=['LEFT(E5,1)="×"'], font=BOLD))
wx.freeze_panes = "C5"; wx.auto_filter.ref = f"A4:L{XL}"; landscape(wx)

# ============ 経歴簿 ============
wc = S["経歴簿"]
title(wc, "潜水経歴簿", None, 4)
VR = lambda c: f"'隊員一覧'!${c}${V_FIRST}:${c}${V_LAST}"
labels = [(4, 2, "職員番号", 3, None, None),
          (4, 5, "氏名", 6, f'=IF($C$4="","",IFERROR(INDEX({M("B")},MATCH($C$4,{M("A")},0)),"番号なし"))', None),
          (4, 8, "所属", 9, f'=IF($C$4="","",IFERROR(INDEX({M("C")},MATCH($C$4,{M("A")},0)),""))', None),
          (5, 2, "潜水士免許", 3, f'=IF($C$4="","",IF({q_cnt("$C$4","潜水士免許")}>0,"有","無"))', None),
          (5, 5, "水難救助研修", 6, f'=IF($C$4="","",IF({q_cnt("$C$4","水難救助研修")}>0,"修了","—"))', None),
          (5, 8, "小型船舶", 9, f'=IF($C$4="","",IFERROR(INDEX({VR("G")},MATCH($C$4,{VR("A")},0)),""))', None),
          (6, 2, "累計本数", 3, f'=IF($C$4="",0,COUNTIF({D("B")},$C$4))', '0"本"'),
          (6, 5, "累計潜水時間", 6, f'=IF($C$4="","",SUMIF({D("B")},$C$4,{D("H")}))', '0"分"'),
          (6, 8, "最大水深", 9, f'=IF(OR($C$4="",$C$6=0),"",_xlfn.MAXIFS({D("G")},{D("B")},$C$4))', '0"m"'),
          (7, 2, "最終潜水日", 3, f'=IF(OR($C$4="",$C$6=0),"",_xlfn.MAXIFS({D("A")},{D("B")},$C$4))', DATE),
          (7, 5, "今年度本数", 6, f"=IF($C$4=\"\",\"\",COUNTIFS({D('B')},$C$4,{D('AF')},YEAR(EDATE(TODAY(),-3))))", '0"本"'),
          (7, 8, "時間換算", 9, '=IF(F6="","",ROUND(F6/60,1))', '0.0"時間"'),
          (2, 8, "出力日", 9, "=TODAY()", DATE)]
for (r_, lc, lab, vc, fm, nf) in labels:
    wc.cell(r_, lc, lab).font = f(9.5)
    if fm:
        c = wc.cell(r_, vc, fm); c.font = f(11, True); c.alignment = Alignment(horizontal="left")
        if nf: c.number_format = nf
c = wc["C4"]; c.value = 1002; c.font = f(12, True); c.border = Border(left=thin, right=thin, top=thin, bottom=thin); c.alignment = Alignment(horizontal="left")
dv_list(wc, "L_職員番号", "C4"); wc["D4"] = "← プルダウンで選ぶ"; wc["D4"].font = f(9)
KH, KF, KN = 9, 10, 300
header(wc, KH, ["No", "日付", "区分", "事案・訓練名", "地点", "最大水深(m)", "潜水時間(分)", "水面休息(分)", "グループ", "減圧停止", "体調異常", "指揮者"])
for col, w in zip("ABCDEFGHIJKL", [5, 11, 12, 22, 12, 11, 9, 9, 12, 8, 8, 11]): wc.column_dimensions[col].width = w
src = {2: "A", 3: "D", 4: "E", 5: "F", 6: "G", 7: "H", 8: "I", 9: "X", 10: "O", 11: "P", 12: "Q"}
for i in range(KN):
    r = KF + i; wc.cell(r, 1, f"=IF({i+1}<=$C$6,{i+1},\"\")")
    for c, s in src.items():
        ix = f'INDEX({D(s)},MATCH($A{r},{D("AG")},0))'; wc.cell(r, c, f'=IF($A{r}="","",IF({ix}="","",{ix}))')
    wc.cell(r, 2).number_format = DATE
body(wc, KF, KF + KN - 1, 1, 12, center=(1, 3, 6, 7, 8, 9, 10, 11))
wc.conditional_formatting.add(f"K{KF}:K{KF+KN-1}", FormulaRule(formula=[f'K{KF}="有"'], font=BOLD))
wc.freeze_panes = f"A{KF}"; landscape(wc); wc.print_title_rows = f"{KH}:{KH}"

# ============ 報告書 ============
wr = S["報告書"]
title(wr, "潜水活動報告書", None, 6)
KX = lambda c: f'IFERROR(INDEX({K(c)},MATCH($C$4,{K("A")},0))&"","")'
def lab(r_, c_, text): x = wr.cell(r_, c_, text); x.font = f(9.5)
def val(r_, c_, fm, nf=None, merge_to=None, size=11):
    if merge_to: wr.merge_cells(start_row=r_, start_column=c_, end_row=r_, end_column=merge_to)
    x = wr.cell(r_, c_, fm); x.font = f(size, True); x.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    if nf: x.number_format = nf
lab(4, 2, "事案ID")
c = wr["C4"]; c.value = list(incs.keys())[-1]; c.font = f(12, True); c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
wr.merge_cells("C4:E4"); dv_list(wr, "L_事案ID", "C4"); lab(4, 6, "← プルダウンで選ぶ")
lab(2, 9, "出力日"); val(2, 10, "=TODAY()", DATE, 12)
lab(5, 2, "日付"); val(5, 3, f'=IF($C$4="","",IFERROR(INDEX({K("B")},MATCH($C$4,{K("A")},0)),""))', DATE, 5)
lab(5, 6, "区分"); val(5, 7, f'=IF($C$4="","",{KX("C")})', None, 8)
lab(5, 9, "時間"); val(5, 10, f'=IF($C$4="","",RIGHT({KX("G")},5)&"〜"&RIGHT({KX("H")},5))', None, 12)
lab(6, 2, "事案名"); val(6, 3, f'=IF($C$4="","",{KX("D")})', None, 8)
lab(6, 9, "指揮者"); val(6, 10, f'=IF($C$4="","",{KX("F")})', None, 12)
lab(7, 2, "地点"); val(7, 3, f'=IF($C$4="","",{KX("E")})', None, 8)
lab(7, 9, "海況"); val(7, 10, f'=IF($C$4="","",{KX("K")})', None, 12)
lab(8, 2, "水温"); val(8, 3, f'=IF($C$4="","",IF({KX("I")}="","—",{KX("I")}&"℃"))', None, 5)
lab(8, 6, "透視度"); val(8, 7, f'=IF($C$4="","",IF({KX("J")}="","—",{KX("J")}&"m"))', None, 8)
lab(8, 9, "潜水表"); val(8, 10, f'=IF($C$4="","",{KX("M")})', None, 12, size=9)
lab(9, 2, "備考"); val(9, 3, f'=IF($C$4="","",{KX("L")})', None, 12)
wr.row_dimensions[9].height = 30
for c_ in range(2, 13): wr.cell(10, c_).border = Border(bottom=med)
lab(11, 2, "潜水員"); val(11, 3, f'=IF($C$4="","",COUNTIFS({D("S")},$C$4,{D("T")},1))', '0"人"', 5)
lab(11, 6, "総本数"); val(11, 7, f'=IF($C$4="",0,COUNTIF({D("S")},$C$4))', '0"本"', 8)
lab(11, 9, "最大水深"); val(11, 10, f'=IF(OR($C$4="",$G$11=0),"",_xlfn.MAXIFS({D("G")},{D("S")},$C$4))', '0"m"', 12)
lab(12, 2, "減圧停止"); val(12, 3, f'=IF($C$4="","",COUNTIFS({D("S")},$C$4,{D("O")},"有"))', '0"本"', 5)
lab(12, 6, "体調異常"); val(12, 7, f'=IF($C$4="","",COUNTIFS({D("S")},$C$4,{D("P")},"有"))', '0"本"', 8)
lab(12, 9, "範囲外"); val(12, 10, f'=IF($C$4="","",COUNTIFS({D("S")},$C$4,{D("AC")},"範囲外"))', '0"本"', 12)
RH, RF, RN = 14, 15, 30
header(wr, RH, ["No", "潜水員", "本数", "潜降", "浮上", "最大水深(m)", "潜水時間(分)", "水面休息(分)", "グループ", "判定", "体調異常", "備考"])
for col, w in zip("ABCDEFGHIJKL", [5, 12, 6, 7, 7, 9, 9, 9, 8, 11, 8, 22]): wr.column_dimensions[col].width = w
rsrc = {2: "C", 3: "T", 4: "U", 5: "V", 6: "G", 7: "H", 8: "I", 9: "X", 10: "AC", 11: "P", 12: "R"}
for i in range(RN):
    r = RF + i; wr.cell(r, 1, f"=IF({i+1}<=$G$11,{i+1},\"\")")
    for c_, s in rsrc.items():
        ix = f'INDEX({D(s)},MATCH($A{r},{D("AJ")},0))'; wr.cell(r, c_, f'=IF($A{r}="","",IF({ix}="","",{ix}))')
    wr.cell(r, 4).number_format = TIME; wr.cell(r, 5).number_format = TIME
body(wr, RF, RF + RN - 1, 1, 12, center=(1, 3, 4, 5, 6, 7, 8, 9, 11))
for r in range(RF, RF + RN): wr.cell(r, 12).alignment = Alignment(wrap_text=True, vertical="center"); wr.cell(r, 12).font = f(9)
wr.conditional_formatting.add(f"J{RF}:K{RF+RN-1}", FormulaRule(formula=[f'OR(LEFT(J{RF},2)="減圧",J{RF}="範囲外",K{RF}="有")'], font=BOLD))
SR = RF + RN + 1
lab(SR, 2, "報告者"); lab(SR, 6, "確認者（隊長）"); lab(SR, 9, "決裁")
for c_ in (3, 4, 7, 8, 10, 11, 12): wr.cell(SR, c_).border = Border(bottom=thin)
wr.row_dimensions[SR].height = 30
wr.cell(SR + 1, 2, "アプリの判定はNAUI潜水表による計算の目安です。実際の判断は指揮者が行いました。").font = f(8.5)
landscape(wr, portrait=True); wr.print_area = f"A1:L{SR+1}"
wr.page_setup.fitToHeight = 1

# ============ CSV抽出 ============
we = S["CSV抽出"]
title(we, "CSV抽出", "条件を入れると下の表に該当する潜水記録が並びます。空欄の条件は「指定なし」。［CSVで保存］で書き出せます。", 4)
for r_, labx, v, nf in [(4, "開始日", dt.date(2026, 4, 1), DATE), (5, "終了日", dt.date(2027, 3, 31), DATE), (6, "区分", "すべて", None), (7, "職員番号", None, None)]:
    we.cell(r_, 2, labx).font = f(9.5); c = we.cell(r_, 3, v); c.font = f(11, True)
    c.border = Border(left=thin, right=thin, top=thin, bottom=thin); c.alignment = Alignment(horizontal="left")
    if nf: c.number_format = nf
dv_list(we, "L_抽出区分", "C6"); dv_list(we, "L_職員番号", "C7")
we.cell(8, 2, "該当件数").font = f(9.5)
c = we.cell(8, 3, f"=COUNTIF({D('AH')},1)"); c.font = f(11, True); c.number_format = '0"件"'; c.alignment = Alignment(horizontal="left")
we.cell(8, 5, "500件を超える場合は期間を分けて抽出してください。").font = f(9)
EH, EF, EN = 10, 11, 500
header(we, EH, ["No"] + DCOLS, [5] + DW)
for i in range(EN):
    r = EF + i; we.cell(r, 1, f"=IF({i+1}<=$C$8,{i+1},\"\")")
    for j in range(31):
        ix = f'INDEX({D(CL(1+j))},MATCH($A{r},{D("AI")},0))'; we.cell(r, 2 + j, f'=IF($A{r}="","",IF({ix}="","",{ix}))')
    we.cell(r, 2).number_format = DATE; we.cell(r, 22).number_format = TIME; we.cell(r, 23).number_format = TIME
body(we, EF, EF + EN - 1, 1, 32)
we.freeze_panes = f"A{EF}"

# ============ 潜水表 ============
wt = S["潜水表"]
title(wt, "潜水表", "アプリに入れる潜水表の正本です。数値を直したら、2人で照合してから［アプリ用に書き出す］を押し、できたファイルを現場用スマホへメールで送ります。", 6)
T = TABLE
LET = list("ABCDEFGHIJKLMN")
meta = [("表の名前", T["name"]), ("版", T["version"]), ("出典", T["source"]), ("照合", "照合済み" if T["verified"] else "照合前"),
        ("照合者", T["verifiedBy"]), ("照合日", T["updatedAt"]), ("備考", "／".join(T["notes"]))]
for i, (k_, v) in enumerate(meta):
    wt.cell(3 + i, 1, k_).font = f(9.5, True); c = wt.cell(3 + i, 2, v); c.font = f()
    wt.merge_cells(start_row=3 + i, start_column=2, end_row=3 + i, end_column=8); c.border = Border(bottom=hair)
dvv = DataValidation(type="list", formula1='"照合済み,照合前"', allow_blank=False); wt.add_data_validation(dvv); dvv.add("B6")
rules = [("最短の水面休息（分）", "minSurfaceInterval"), ("1本目に戻る休息（分）", "clearAfterMinutes"), ("反復を避ける水深（m超）", "avoidRepetitiveDeeperThan"),
         ("安全停止の時間（分）", "safetyStopMinutes"), ("安全停止の水深（m）", "safetyStopDepth"), ("浮上速度の上限（m/分）", "maxAscentRate")]
for i, (k_, key) in enumerate(rules):
    wt.cell(3 + i, 10, k_).font = f(9.5, True); c = wt.cell(3 + i, 13, T["rules"][key]); c.font = f(); c.border = Border(bottom=hair)
    wt.merge_cells(start_row=3 + i, start_column=10, end_row=3 + i, end_column=12)
GRIDS = {}
def grid(top, name, note, row_head, row_keys, col_keys, getter, col_title, fmt=None):
    wt.cell(top, 1, name).font = f(12, True); wt.cell(top, 4, note).font = f(9)
    h = top + 1
    wt.cell(h, 1, row_head).font = f(9.5, True)
    for j, ck in enumerate(col_keys): c = wt.cell(h, 2 + j, col_title(ck)); c.font = f(9.5, True); c.alignment = Alignment(horizontal="center")
    for c_ in range(1, 2 + len(col_keys)): wt.cell(h, c_).border = Border(bottom=med)
    for i, rk in enumerate(row_keys):
        r = h + 1 + i; c = wt.cell(r, 1, rk); c.font = f(10, True); c.alignment = Alignment(horizontal="center")
        for j, ck in enumerate(col_keys):
            v = getter(rk, ck); x = wt.cell(r, 2 + j, v); x.font = f(); x.alignment = Alignment(horizontal="center")
            x.border = Border(left=hair, right=hair, top=hair, bottom=hair)
            if fmt: x.number_format = fmt
    GRIDS[name] = (h, len(row_keys), len(col_keys))
    return h + len(row_keys) + 3
depths = T["depths"]
gt = {d_: {g: tm for tm, g in T["groups"][str(d_)]} for d_ in depths}
hm_ = lambda m: f"{m//60}:{m%60:02d}"
si_lo = {(ng, eg): lo for eg, rows in T["surfaceInterval"].items() for lo, hi, ng in rows}
si_hi = {(ng, eg): hi for eg, rows in T["surfaceInterval"].items() for lo, hi, ng in rows}
top = 11
wt.cell(top, 1, "最大潜水時間（分）").font = f(12, True)
wt.cell(top + 1, 1, "水深(m)").font = f(9.5, True)
for j, d_ in enumerate(depths):
    wt.cell(top + 1, 2 + j, d_).font = f(9.5, True); x = wt.cell(top + 2, 2 + j, T["mdt"][str(d_)]); x.font = f(); x.border = Border(left=hair, right=hair, top=hair, bottom=hair)
    wt.cell(top + 1, 2 + j).alignment = Alignment(horizontal="center"); x.alignment = Alignment(horizontal="center")
wt.cell(top + 2, 1, "MDT").font = f(10, True)
for c_ in range(1, 2 + len(depths)): wt.cell(top + 1, c_).border = Border(bottom=med)
GRIDS["MDT"] = (top + 1, 1, len(depths))
top += 5
top = grid(top, "表1 グループの時間（分）", "その時間までの潜水で、列のグループになります。", "水深(m)", depths, LET, lambda r, c: gt[r].get(c), lambda c: c)
top = grid(top, "表1 減圧停止（分）", "上の表で減圧が必要な欄に、5mでの停止時間を入れます。", "水深(m)", depths, LET,
           lambda r, c: T["deco"].get(f"{r}-{c}", {}).get("stop5"), lambda c: c)
top = grid(top, "表3 残留窒素時間（RNT・分）", "", "グループ", LET, depths, lambda r, c: T["rnt"].get(r, {}).get(str(c)), lambda c: f"{c}m")
top = grid(top, "表3 修正最大潜水時間（AMDT・分）", "空欄は「表に記載なし」（潜らない扱い）。", "グループ", LET, depths, lambda r, c: T["amdt"].get(r, {}).get(str(c)), lambda c: f"{c}m")
top = grid(top, "表2 水面休息 下限（時:分）", "行＝新しいグループ、列＝潜水終了時のグループ。", "新グループ", LET, LET, lambda r, c: hm_(si_lo[(r, c)]) if (r, c) in si_lo else None, lambda c: f"終了{c}", "@")
top = grid(top, "表2 水面休息 上限（時:分）", "", "新グループ", LET, LET, lambda r, c: hm_(si_hi[(r, c)]) if (r, c) in si_hi else None, lambda c: f"終了{c}", "@")
wt.column_dimensions["A"].width = 11
for j in range(2, 17): wt.column_dimensions[CL(j)].width = 6.5
wt.freeze_panes = "B3"
json.dump(GRIDS, open(BUILD/'grids.json', 'w'), ensure_ascii=False)

# ============ 使い方 ============
title(ws0, "潜水管理台帳 第2版", "宮古島市消防本部 潜水隊。Excel 2019対応。一覧と集計は数式だけで動き、マクロは取り込み・書き出し用。氏名・数値はすべて見本です。", 6)
rows = [("", ""), ("見る（自動）", ""),
        ("隊員一覧", "1人1行。免許・研修・小型船舶・潜水本数・ブランクを一覧表示。太字は要確認"),
        ("資格一覧", "隊員×資格の○—表"),
        ("経歴簿", "職員番号を選ぶと、その人の潜水歴と累計が並ぶ。PDFで出せる"),
        ("報告書", "事案IDを選ぶと、潜水活動報告書が1枚にまとまる。PDFで出せる"),
        ("CSV抽出", "期間・区分・職員番号で潜水記録を絞り込み、CSVで出す"),
        ("", ""), ("入力する", ""),
        ("潜水記録・事案", "アプリのCSVを取り込むと自動で入る。手で入力もできる"),
        ("隊員マスタ", "1人1行。職員番号は重複不可。異動・退職しても行は消さず在籍状態を変える"),
        ("資格研修記録", "1件1行。小型船舶は級ごとに1行（1級と特殊なら2行）"),
        ("潜水表", "アプリの潜水表の正本。直したら2人で照合し［アプリ用に書き出す］"),
        ("選択肢", "プルダウンの中身"),
        ("", ""), ("アプリの記録を取り込む", ""),
        ("1", "現場用スマホから届いたメールのCSVを、パソコンに保存する"),
        ("2", "潜水記録シートの［アプリのCSVを取り込む］→ CSVを選ぶ（複数選べる）"),
        ("3", "事案と潜水記録に追加される。同じ記録（事案ID・職員番号・本数が同じ）は追加されない"),
        ("4", "報告書シートで事案IDを選び［報告書をPDF］"),
        ("", ""), ("潜水表をアプリに送る", ""),
        ("1", "潜水表シートの数値を直し、2人で照合して「照合」欄を照合済みにする"),
        ("2", "［アプリ用に書き出す］→「出力」フォルダに .json ができる"),
        ("3", "その .json を現場用スマホへメールで送り、アプリの 潜水表 →［潜水表ファイルを読み込む］"),
        ("", ""), ("マクロの組み込み（初回だけ）", ""),
        ("1", "このファイルを［名前を付けて保存］→ 種類「Excel マクロ有効ブック（.xlsm）」"),
        ("2", "Alt＋F11 →［ファイル］→［ファイルのインポート］→ 潜水管理マクロ.bas"),
        ("3", "Alt＋F8 →「SetupButtons」を実行。各シートにボタンが付く"),
        ("4", "上書き保存。次回から開くときに［コンテンツの有効化］"),
        ("", ""), ("注意", ""),
        ("判定", "アプリの判定はNAUI潜水表による計算の目安。実際の判断は指揮者とダイブコンピュータを優先"),
        ("上限", "隊員200人、資格研修記録1000件、潜水記録3000件、事案1000件まで")]
ws0.column_dimensions["A"].width = 22; ws0.column_dimensions["B"].width = 100
for i, (a, b) in enumerate(rows):
    r = 3 + i; ca = ws0.cell(r, 1, a); cb = ws0.cell(r, 2, b)
    if a and not b:
        ca.font = f(12, True); ca.border = Border(bottom=med); ws0.cell(r, 2).border = Border(bottom=med); ws0.row_dimensions[r].height = 24
    else:
        ca.font = f(10.5, True); cb.font = f(10.5); cb.alignment = Alignment(wrap_text=True, vertical="top"); ca.alignment = Alignment(vertical="top")

for ws in wb.worksheets: ws.sheet_view.showGridLines = ws.title in ("隊員マスタ", "資格研修記録", "潜水記録", "事案", "選択肢")
wb.active = 1
wb.save(OUT); print("saved", OUT)

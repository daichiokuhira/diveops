"""台帳の取り込み（マクロ ImportAppCsv と同じ手順の Python 版 ledger/import_port.py）の試験。
アプリのCSV（48列）と旧形式（35列）を作って取り込み、照合・更新・削除・記録更新日時の順序・職員番号・先頭の ' ・数式を壊さないことを確かめる"""
import subprocess, sys, json, pathlib, tempfile, copy, re
ROOT = pathlib.Path(__file__).resolve().parents[1]
B = ROOT / 'ledger' / 'build'; B.mkdir(exist_ok=True)
subprocess.run(f'node tests/sample_rows.js > "{B/"sample_rows.json"}"', shell=True, check=True, cwd=ROOT)
SAMPLE = json.load(open(B / 'sample_rows.json', encoding='utf-8'))
HEAD = list(SAMPLE[0].keys())
OLD35 = HEAD[:35]; OLD35[14] = '減圧停止'   # 旧形式：35列、15列目の名前は「減圧停止」、記録IDなし

fails = []
def ok(cond, label, detail=''):
    print(('OK  ' if cond else 'NG  ') + label + (f'  {detail}' if detail else ''))
    if not cond: fails.append(label)

def q(v):
    # アプリ（app.js の buildCSV）と同じ書き方
    s = '' if v is None else str(v)
    if isinstance(v, str) and re.match(r'^[=+\-@\t]', s): s = "'" + s
    return f'"{s.replace(chr(34), chr(34) * 2)}"' if re.search(r'[",\r\n]', s) else s

def write_csv(path, rows, head=HEAD, enc='utf-8'):
    lines = [','.join(q(h) for h in head)]
    for r in rows: lines.append(','.join(q(r.get('減圧停止(表の判定)') if h == '減圧停止' else r.get(h, '')) for h in head))
    text = ('﻿' if enc == 'utf-8' else '') + '\r\n'.join(lines) + '\r\n'
    open(path, 'wb').write(text.encode(enc))
    return path

with tempfile.TemporaryDirectory() as td:
    td = pathlib.Path(td)
    x = td / 'ledger.xlsx'
    subprocess.run([sys.executable, 'ledger/build_ledger.py', str(x)], check=True, cwd=ROOT, capture_output=True)
    sys.path.insert(0, str(ROOT / 'ledger'))
    import import_port as P
    from openpyxl import load_workbook
    wb = load_workbook(x); wd = wb['潜水記録']; wk = wb['事案']
    D1, D2 = P.D_FIRST, P.D_LAST
    snap = lambda: {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None}
    before = snap()
    formulas_before = {k: v for k, v in before.items() if isinstance(v, str) and v.startswith('=')}
    n_rows = lambda: sum(1 for r in range(D1, D2 + 1) if wd.cell(r, 2).value not in (None, ''))
    find = lambda rid: [r for r in range(D1, D2 + 1) if wd.cell(r, P.C_RECID).value == rid]
    start = n_rows()

    # 見本（台帳を作ったときに入れたもの）をそのまま取り込む → 全部「取り込み済み」（台帳を作る処理と取り込みの読み方がそろっている）
    n = P.import_csv(wb, [write_csv(td / 'sample.csv', SAMPLE)])
    ok(n['added'] == 0 and n['updated'] == 0 and n['skipped'] == len(SAMPLE), '見本CSVは全部 取り込み済み', str({k: n[k] for k in ('added', 'updated', 'skipped')}))

    # 新しい事案（職員番号 0123・先頭 ' の付く値を含む）
    new = []
    for i, base in enumerate(SAMPLE[:3]):
        r = copy.deepcopy(base)
        r.update({'事案ID': 'S-20261024', '日付': '2026/10/24', '記録ID': f'rec-test-{i}', '事案・訓練名': '取り込み試験', '本数': 1, '記録更新日時': '2026/10/24 10:00:00'})
        new.append(r)
    new[0]['職員番号'] = '0123'; new[0]['氏名'] = '見本 九郎'
    new[1]['備考'] = '=1+2 と書いたメモ'; new[2]['備考'] = '-3m付近で停止'; new[2]['監視員'] = '@見本'
    new[1]['職員番号'] = 'A1002'
    new[1]['水温(℃)'] = '-1'           # 数値の列でも '- は取り除く
    new[0]['備考'] = "'見本のメモ"      # 次の文字が = + - @ タブでなければ ' は残す
    f1 = write_csv(td / 'new.csv', new)
    # (a) 2回取り込む
    n1 = P.import_csv(wb, [f1]); n2 = P.import_csv(wb, [f1])
    ok(n1['added'] == 3 and n1['inc_added'] == 1, '(a) 1回目は3件追加・事案1件', str(n1))
    ok(n2['added'] == 0 and n2['skipped'] == 3 and n2['inc_added'] == 0, '(a) 2回目は追加0', str(n2))
    # (d) 職員番号 0123 は文字のまま。2回目も二重にならない
    r0 = find('rec-test-0')
    ok(len(r0) == 1 and wd.cell(r0[0], 2).value == '0123' and wd.cell(r0[0], 2).number_format == '@', '(d) 職員番号 0123 が文字のまま1行', repr([wd.cell(r, 2).value for r in r0]))
    ok(len(find('rec-test-1')) == 1 and wd.cell(find('rec-test-1')[0], 2).value == 'A1002', '(d) 英字を含む職員番号')
    # (e) 先頭の ' を1文字取り除く（文字の列だけ）。セルは文字のまま（数式にならない）
    c1 = wd.cell(find('rec-test-1')[0], 18); c2 = wd.cell(find('rec-test-2')[0], 18)
    ok(c1.value == '=1+2 と書いたメモ' and c1.data_type == 's', "(e) '= の ' を取り除き、文字として入る", repr(c1.value))
    ok(c2.value == '-3m付近で停止' and wd.cell(find('rec-test-2')[0], 46).value == '@見本', "(e) '- と '@ も取り除く")
    ok(wd.cell(find('rec-test-1')[0], 13).value == -1, "(e) 数値の列の '-1 は -1（数値）", repr(wd.cell(find('rec-test-1')[0], 13).value))
    ok(wd.cell(find('rec-test-0')[0], 18).value == "'見本のメモ", "(e) 後ろが記号でない ' は残す", repr(wd.cell(find('rec-test-0')[0], 18).value))
    # 新しい列（AL〜AV）
    rr = find('rec-test-0')[0]
    got = {h: wd.cell(rr, c).value for h, c, k in P.MAP if c >= P.C_RECID}
    ok(got['予定水深(m)'] == new[0]['予定水深(m)'] and str(got['浮上開始時刻'])[:5] == new[0]['浮上開始時刻'] and got['呼吸ガス'] == '空気'
       and got['記録更新日時'] == '2026/10/24 10:00:00' and P.C_TS == 49, '新しい列が AL 以降に入る（記録更新日時は AW）', str(got))
    ok(wd.cell(rr, 15).value == new[0]['減圧停止(表の判定)'], '15列目「減圧停止(表の判定)」')

    # (b) 中身を変え、記録更新日時が新しい → 更新
    chg = copy.deepcopy(new); chg[0]['終了圧(気圧)'] = 55; chg[0]['修正記録'] = '1回・最終 2026/10/24 11:00'; chg[0]['記録更新日時'] = '2026/10/24 11:00:00'
    n3 = P.import_csv(wb, [write_csv(td / 'chg.csv', chg)])
    ok(n3['updated'] == 1 and n3['skipped'] == 2 and n3['added'] == 0, '(b) 新しくなった1件は更新', str(n3))
    ok(wd.cell(find('rec-test-0')[0], 12).value == 55 and len(find('rec-test-0')) == 1, '(b) 同じ行が上書きされる')
    stale = copy.deepcopy(new); stale[0]['終了圧(気圧)'] = 77   # 中身は違うが記録更新日時が古い
    n3b = P.import_csv(wb, [write_csv(td / 'stale.csv', stale)])
    ok(n3b['updated'] == 0 and wd.cell(find('rec-test-0')[0], 12).value == 55, '(b) 記録更新日時が古いCSVでは上書きしない', str(n3b))

    # (c) 状態=削除（記録更新日時が新しい）→ 行が空き、記録IDと記録更新日時だけ残る。もう一度送っても何も起きない
    dele = copy.deepcopy(chg); dele[2]['状態'] = '削除'; dele[2]['記録更新日時'] = '2026/10/24 12:00:00'
    row2 = find('rec-test-2')[0]
    n4 = P.import_csv(wb, [write_csv(td / 'del.csv', dele)])
    ok(n4['deleted'] == 1 and n4['added'] == 0, '(c) 削除1件', str(n4))
    keep = {P.C_RECID, P.C_TS, P.C_NOTE}
    ok(all(wd.cell(row2, c).value is None for h, c, k in P.MAP if c not in keep) and str(wd.cell(row2, 3).value).startswith('=')
       and wd.cell(row2, P.C_RECID).value == 'rec-test-2' and wd.cell(row2, P.C_TS).value == '2026/10/24 12:00:00', '(c) 入力列が空き、数式と削除の印（記録ID・記録更新日時）が残る')
    n5 = P.import_csv(wb, [td / 'del.csv'])
    ok(n5['deleted'] == 0 and n5['added'] == 0 and n5['updated'] == 0 and n5['skipped'] == 3, '(c) 2回目は何もしない', str(n5))
    # 未登録の記録の削除が先に届く → 削除の印の行を作る。後から古い v1 を取り込んでも追加しない
    v1u = copy.deepcopy(new[:1]); v1u[0].update({'事案ID': 'S-20261025', '記録ID': 'rec-never', '本数': 9, '記録更新日時': '2026/10/25 09:00:00'})
    unreg = copy.deepcopy(v1u); unreg[0].update({'状態': '削除', '記録更新日時': '2026/10/25 10:00:00'})
    n6 = P.import_csv(wb, [write_csv(td / 'del2.csv', unreg)])
    tomb = find('rec-never')
    ok(n6['deleted'] == 1 and n6['added'] == 0 and n6['inc_added'] == 0 and len(tomb) == 1, '(c) 未登録の削除は削除の印の行を作る（削除1件）', str(n6))
    ok(tomb and {c for h, c, k in P.MAP if wd.cell(tomb[0], c).value is not None} == {P.C_RECID, P.C_TS, P.C_NOTE}
       and wd.cell(tomb[0], P.C_NOTE).value == '削除済み（アプリで削除）', '(c) 印の行は 記録ID・記録更新日時・削除済み だけ')
    n6b = P.import_csv(wb, [write_csv(td / 'v1u.csv', v1u)])
    ok(n6b['added'] == 0 and n6b['inc_added'] == 0 and find('rec-never') == tomb and wd.cell(tomb[0], 2).value is None, '(c) 削除が先・古い v1 が後 → 追加0', str(n6b))

    # 順序：v2（更新・削除）の後に v1 を取り込んでも、戻らない・消した記録が復活しない
    n_v1 = P.import_csv(wb, [f1])
    ok(n_v1['added'] == 0 and n_v1['updated'] == 0 and n_v1['skipped'] == 3, '(順序) v2 の後に v1 を取り込んでも何も変わらない', str(n_v1))
    ok(wd.cell(find('rec-test-0')[0], 12).value == 55 and find('rec-test-2') == [row2] and wd.cell(row2, 2).value is None, '(順序) 更新が戻らず、削除した記録も復活しない')
    # 署で手で直した内容は、同じCSVの取り込み直しで消えない
    r1 = find('rec-test-1')[0]; wd.cell(r1, 18).value = '署で手で直した'
    n_h = P.import_csv(wb, [td / 'del.csv'])
    ok(n_h['updated'] == 0 and wd.cell(r1, 18).value == '署で手で直した', '(手直し) 同じCSVを取り込み直しても手で直した内容が残る', str(n_h))
    # もっと新しいCSVは更新する
    v3 = copy.deepcopy(dele); v3[1]['備考'] = 'アプリで直した'; v3[1]['記録更新日時'] = '2026/10/24 13:00:00'
    n_v3 = P.import_csv(wb, [write_csv(td / 'v3.csv', v3)])
    ok(n_v3['updated'] == 1 and wd.cell(r1, 18).value == 'アプリで直した', '(順序) もっと新しいCSVは更新する', str(n_v3))

    # (f) 旧形式 35列
    old = [dict(r, **{'事案ID': 'S-20261001', '日付': '2026/10/01', '事案・訓練名': '旧形式の試験'}) for r in SAMPLE[:2]]
    fo = write_csv(td / 'old.csv', old, OLD35)
    n7 = P.import_csv(wb, [fo]); n8 = P.import_csv(wb, [fo])
    ok(n7['added'] == 2 and n8['added'] == 0 and n8['skipped'] == 2, '(f) 旧形式も取り込め、2回目は追加0', f'{n7} {n8}')
    ro = [r for r in range(D1, D2 + 1) if wd.cell(r, 19).value == 'S-20261001']
    ok(len(ro) == 2 and wd.cell(ro[0], 15).value == old[0]['減圧停止(表の判定)'] and wd.cell(ro[0], P.C_RECID).value is None, '(f) 「減圧停止」が15列目に入る')
    # 旧形式で入れた行に、記録IDつきの新しいCSVが来た → 二重にせず、その行に記録IDを入れる（更新）
    newer = [dict(r, **{'記録ID': f'rec-old-{i}'}) for i, r in enumerate(old)]
    n9 = P.import_csv(wb, [write_csv(td / 'newer.csv', newer)])
    ok(n9['added'] == 0 and n9['updated'] == 2 and wd.cell(ro[0], P.C_RECID).value == 'rec-old-0', '(f) 旧形式の行に記録IDが付く（二重にならない）', str(n9))

    # 次に書く行：日付が空でも職員番号・事案IDのある行には上書きしない
    last = P.last_data_row(wd, D1, D2, P.NEXT_COLS)
    wd.cell(last + 1, 2).value = '1001'; wd.cell(last + 1, 19).value = '手入力'; wd.cell(last + 1, 20).value = 1
    extra = [dict(SAMPLE[0], **{'事案ID': 'S-20261002', '記録ID': 'rec-next', '本数': 1})]
    P.import_csv(wb, [write_csv(td / 'next.csv', extra)])
    ok(find('rec-next') == [last + 2] and wd.cell(last + 1, 19).value == '手入力', '(6) 日付が空の行を飛ばして次の行に書く', f'{find("rec-next")} {last + 2}')

    # (5) Excel で上書き保存した（Shift_JIS の）CSV は取り込まない
    nsj = P.import_csv(wb, [write_csv(td / 'sjis.csv', [dict(SAMPLE[0], **{'記録ID': 'rec-sjis'})], enc='cp932')])
    ok(nsj['added'] == 0 and len(nsj['notes']) == 1 and 'Excelで保存し直したCSV' in nsj['notes'][0], '(5) Shift_JIS のCSVは知らせて取り込まない', nsj['notes'][0] if nsj['notes'] else '')
    nbad = P.import_csv(wb, [write_csv(td / 'other.csv', [{'a': 1}], ['a', 'b'])])
    ok(nbad['added'] == 0 and 'アプリのCSVではありません' in nbad['notes'][0], 'アプリのCSVでないファイルは知らせる')

    # (g) 数式を壊していない：取り込みの前後で数式は1つも変わらず、書き換えたのは入力の列だけ
    out = td / 'after.xlsx'; wb.save(out)
    wb2 = load_workbook(out); after = {(ws.title, c.coordinate): c.value for ws in wb2.worksheets for row in ws.iter_rows() for c in row if c.value is not None}
    formulas_after = {k: v for k, v in after.items() if isinstance(v, str) and v.startswith('=') and wb2[k[0]][k[1]].data_type == 'f'}
    ok(formulas_after == formulas_before, '(g) 数式の数と中身が取り込み前と同じ', f'{len(formulas_before)}個')
    from openpyxl.utils import column_index_from_string as CI
    from openpyxl.utils.cell import coordinate_from_string as CS
    dive_cols = {c for h, c, k in P.MAP}; inc_cols = {c for h, c, k in P.INC} | {16}
    touched = {k for k in set(before) | set(after) if before.get(k) != after.get(k)}
    bad = [k for k in touched if not ((k[0] == '潜水記録' and CI(CS(k[1])[0]) in dive_cols and D1 <= CS(k[1])[1] <= D2) or
                                      (k[0] == '事案' and CI(CS(k[1])[0]) in inc_cols and P.K_FIRST <= CS(k[1])[1] <= P.K_LAST))]
    ok(not bad, '(g) 入力の列以外は書き換えていない', f'変わったセル {len(touched)}個、範囲外 {bad[:5]}')
    ok(not (dive_cols & {3, 32, 33, 34, 35, 36}), '(g) 取り込みの列と数式の列（C・AF〜AJ）が重ならない')
    # 削除の印の行（日付・職員番号・事案ID が空）は、経歴簿・隊員一覧・事案の本数・CSV抽出・報告書に数えられない
    #  ・ほかのシートが潜水記録を集計・照合する列（INDEX で値を取り出す列を除く）は、印の行ではすべて空か、空の判定で "" / 0 を返す数式
    #  ・印の行に値がある列（AL・AV・AW）は、CSV抽出の INDEX（補助列の MATCH で行を選んだあと）でしか参照されない
    from openpyxl.utils import get_column_letter as CLx
    wd2 = wb2['潜水記録']; tr = [r for r in range(D1, D2 + 1) if wd2.cell(r, P.C_RECID).value == 'rec-never'][0]
    used, idx_only = set(), set()
    for (sh, coord), f in formulas_after.items():
        if sh == '潜水記録': continue
        for col in re.findall(r"'潜水記録'!\$([A-Z]+)\$\d+:\$[A-Z]+\$\d+", re.sub(r"INDEX\('潜水記録'!\$[A-Z]+\$\d+:\$[A-Z]+\$\d+,", 'INDEX(', f)): used.add(col)
        for col in re.findall(r"INDEX\('潜水記録'!\$([A-Z]+)\$", f): idx_only.add(col)
    guards = {3: f'=IF(B{tr}=""', 32: f'=IF(A{tr}=""', 33: f'=IF(OR(B{tr}=""', 34: f'=IF(A{tr}="",0', 35: f'=IF(AH{tr}=1', 36: f'=IF(OR(S{tr}=""'}
    badc = []
    for col in sorted(used):
        c = CI(col); v = wd2.cell(tr, c).value
        if c in guards: (str(v).startswith(guards[c]) or badc.append(col))
        elif v is not None: badc.append(col)
    marked = {CLx(P.C_RECID), CLx(P.C_TS), CLx(P.C_NOTE)}
    ok(not badc and not (marked & used), '(削除の印) 集計・照合に使う列は、印の行では空か空の判定つき', f'集計の列 {sorted(used)} / 問題 {badc}')
    ok(wd2.cell(tr, 1).value is None and wd2.cell(tr, 2).value is None and wd2.cell(tr, 19).value is None, '(削除の印) 日付・職員番号・事案IDが空（本数・経歴簿・報告書・CSV抽出の対象外）')
    ok(wb2['潜水記録'].auto_filter.ref == f"A4:{P.G['layout']['D_LASTCOL']}{D2}", '(7) オートフィルタが最後の列まで', wb2['潜水記録'].auto_filter.ref)
    ok(n_rows() == start + 3 - 1 + 2 + 1 + 1, '行数の合計', f'{start} → {n_rows()}')

    # VBA と Python 版が1対1か（手続きの名前）
    bas = subprocess.run([sys.executable, 'ledger/make_bas.py', str(td / 'm.bas')], check=True, cwd=ROOT, capture_output=True)
    vba = (td / 'm.bas').read_bytes().decode('cp932')
    pairs = [('ImportOne', 'import_one'), ('HeaderMap', 'header_map'), ('IsAppHeader', 'is_app_header'), ('RowKey', 'row_key'), ('Reindex', 'reindex'),
             ('Unindex', 'unindex'), ('FindRow', 'find_row'), ('WriteRow', 'write_row'), ('RowSame', 'row_same'), ('ClearRow', 'clear_row'),
             ('Fld', 'fld'), ('CellStr', 'cell_str'), ('PutVal', 'put_val'), ('SameVal', 'same_val'), ('LastDataRow', 'last_data_row'),
             ('NextEmptyRow', 'next_empty_row'), ('ParseCsv', 'parse_csv'), ('ReadText', 'read_text')]
    miss = [v for v, p in pairs if not re.search(rf'(Sub|Function) {v}\(', vba) or not hasattr(P, p)]
    ok(not miss, 'VBA の手続きと Python 版が対になっている', str(miss))

print('取り込み試験', 'すべて合格' if not fails else f'不合格 {len(fails)}件')
sys.exit(1 if fails else 0)

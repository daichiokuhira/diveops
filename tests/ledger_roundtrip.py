"""台帳の「潜水表」シート → アプリ用ファイル の往復で、数値が1つも変わらないことを確かめる。
あわせて、マクロと同じ読み方（時刻型のセル・コロンのない「10」・日付型の照合日・空の規則）でも止まらず、同じ結果になることを確かめる"""
import subprocess, sys, json, pathlib, tempfile, datetime as dt
ROOT = pathlib.Path(__file__).resolve().parents[1]
B = ROOT / 'ledger' / 'build'; B.mkdir(exist_ok=True)
subprocess.run(f'node tests/sample_rows.js > "{B/"sample_rows.json"}"', shell=True, check=True, cwd=ROOT)

def check(j):
    out = subprocess.run(['node', '-e', f"""
const E=require('./app/src/engine.js');const a=require('./app/tables/naui-2022.json');const b=require({json.dumps(str(j))});
const v=E.validate(b);const d=E.diff(a,b);
console.log(JSON.stringify({{errors:v.errors.length,warnings:v.warnings.length,diff:d.length}}));"""], check=True, cwd=ROOT, capture_output=True, text=True)
    return json.loads(out.stdout)

ok = True
with tempfile.TemporaryDirectory() as td:
    x = pathlib.Path(td) / 'ledger.xlsx'; j = pathlib.Path(td) / 'table.json'
    subprocess.run([sys.executable, 'ledger/build_ledger.py', str(x)], check=True, cwd=ROOT, capture_output=True)
    subprocess.run([sys.executable, 'ledger/export_port.py', str(x), str(j)], check=True, cwd=ROOT)
    r = check(j); print('往復テスト', r); ok &= r == {'errors': 0, 'warnings': 0, 'diff': 0}

    # 手で直された表：時刻型のセル、「0:10」を「10」と入力、照合日が日付型
    from openpyxl import load_workbook
    sys.path.insert(0, str(ROOT / 'ledger')); import export_port
    wb = load_workbook(x); ws = wb['潜水表']
    top = export_port.GR['表2 水面休息 下限（時:分）'][0]
    changed = 0
    for rr in range(top + 1, top + 15):
        for cc in range(2, 16):
            v = ws.cell(rr, cc).value
            if v == '0:10' and changed == 0: ws.cell(rr, cc).value = '10'; changed += 1
            elif isinstance(v, str) and ':' in v and changed in (1, 2):
                h, m = map(int, v.split(':'))
                if h < 24: ws.cell(rr, cc).value = dt.time(h, m); ws.cell(rr, cc).number_format = 'h:mm'; changed += 1
    ws['B8'] = dt.datetime(2026, 10, 24); ws['B8'].number_format = 'yyyy/m/d'
    x2 = pathlib.Path(td) / 'edited.xlsx'; wb.save(x2)
    t2 = export_port.main(str(x2), str(j))
    r = check(j); print('手で直した表の往復', r, '直したセル', changed, '照合日', t2['updatedAt'])
    ok &= r == {'errors': 0, 'warnings': 0, 'diff': 0} and changed == 3 and t2['updatedAt'] == '2026-10-24'

    # 空の規則は 0（VBA の JNum と同じ）。0.5 は 0.5（Str$ の " .5" にならない）
    ws['M8'] = None; ws['M7'] = 0.5; wb.save(x2)
    t3 = export_port.main(str(x2), str(j))
    print('規則', t3['rules']); ok &= t3['rules']['maxAscentRate'] == 0 and t3['rules']['safetyStopDepth'] == 0.5

sys.exit(0 if ok else 1)
